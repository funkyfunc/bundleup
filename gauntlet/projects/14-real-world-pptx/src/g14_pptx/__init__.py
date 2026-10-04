"""Gauntlet 14: python-pptx with an embedded image."""
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


def main() -> int:
    prs = Presentation()  # reads pptx/templates/default.pptx via __file__
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Gauntlet"

    png = BytesIO()
    Image.new("RGB", (32, 32), (200, 30, 30)).save(png, format="PNG")
    png.seek(0)
    slide.shapes.add_picture(png, Inches(1), Inches(1))

    out = BytesIO()
    prs.save(out)
    out.seek(0)
    again = Presentation(out)
    assert again.slides[0].shapes.title.text == "Gauntlet"
    print("GAUNTLET OK 14-real-world-pptx")
    return 0
