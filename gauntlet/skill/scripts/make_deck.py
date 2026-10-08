"""Make a one-slide deck, as a skill's script would. Prints SKILL OK and what it used."""

import os
import sys
import tempfile

import pptx
from helper import title  # beside this script, as python runs it

if sys.platform == "win32":
    import win32api  # pywin32: its .pth adds win32/ to sys.path

    extra = f"win32api {win32api.GetVersion() is not None}"
else:
    extra = "no win32"
deck = pptx.Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[0])
slide.shapes.title.text = title()
out = os.path.join(tempfile.mkdtemp(), "deck.pptx")
deck.save(out)
print("SKILL OK", sys.argv[1:], os.path.getsize(out) > 0, extra)
